"""Structured Word export for Presence Coach developmental reviews."""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


SECTION_HEADINGS = (
    "WHAT THE COACH DID WELL",
    "MARKER / BEHAVIORAL EVIDENCE",
    "COMPETENCY SYNTHESIS",
    "PATTERNS TO WATCH",
    "THREE HIGH-LEVERAGE PRACTICE EDGES",
    "MOMENTS WORTH REVISITING",
    "BOTTOM LINE",
)

EVIDENCE_STATUSES = (
    "OBSERVED",
    "PARTIAL EVIDENCE",
    "NOT OBSERVED",
    "NO OPPORTUNITY",
    "NOT ASSESSABLE",
    "LIMITED EVIDENCE",
    "CONSISTENT",
    "INCONSISTENT",
    "EXCEEDS THE STANDARD",
    "MEETS THE STANDARD",
    "BELOW THE STANDARD",
    "DOES NOT MEET STANDARD",
    "N/A",
)

_STATUS_ALTERNATION = "|".join(re.escape(item) for item in EVIDENCE_STATUSES)
STATUS_PATTERN = re.compile(
    rf"^(?P<name>.+?)\s+[—-]\s+(?P<status>{_STATUS_ALTERNATION})\s*$",
    re.IGNORECASE,
)

COMPETENCY_PATTERN = re.compile(
    r"^(?P<name>.+?)\s+[—-]\s+Evidence strength:\s*(?P<strength>.+?)\s*$",
    re.IGNORECASE,
)

TIMESTAMP_PATTERN = re.compile(r"(?<!\d)(\d{2}:\d{2}:\d{2})(?!\d)")


def _clean(value: str) -> str:
    return re.sub(r"\s+", " ", (value or "").strip())


def _strip_markdown(value: str, *, strip_bullet: bool = False) -> str:
    """Normalize common Gemini/Groq markdown without losing the actual review text."""
    line = (value or "").strip()
    if not line:
        return ""

    line = re.sub(r"^#{1,6}\s*", "", line)
    if re.match(r"^\*\s+", line):
        line = re.sub(r"^\*\s+", "- ", line)

    line = line.replace("**", "").replace("__", "").replace("`", "")
    line = line.replace("***", "").strip()
    line = re.sub(r"^\*+(?=\S)", "", line)
    line = re.sub(r"\*+$", "", line).strip()

    if strip_bullet:
        line = re.sub(r"^[-+•]\s*", "", line).strip()
    return line


def _display_name(value: str) -> str:
    name = _clean(value)
    if name.startswith("[") and name.endswith("]"):
        name = name[1:-1].strip()
    return name


def split_review_sections(review_text: str) -> dict[str, list[str]]:
    """Split review text even when the provider returns markdown headings/emphasis."""
    sections: dict[str, list[str]] = {"PREAMBLE": []}
    current = "PREAMBLE"

    for raw in (review_text or "").splitlines():
        probe = _strip_markdown(raw, strip_bullet=True)
        probe_upper = probe.upper()

        if probe_upper in SECTION_HEADINGS:
            current = probe_upper
            sections.setdefault(current, [])
            continue

        if probe_upper.startswith("DEVELOPMENTAL REVIEW"):
            sections["TITLE"] = [probe]
            continue

        if probe_upper.startswith("ASSESSMENT BASIS:"):
            sections["ASSESSMENT BASIS"] = [probe.split(":", 1)[1].strip()]
            continue

        normalized = _strip_markdown(raw, strip_bullet=False)
        sections.setdefault(current, []).append(normalized)

    return sections


def parse_behavior_evidence(lines: list[str]) -> list[dict[str, str]]:
    """Parse behavior blocks from plain text or markdown-formatted AI output."""
    rows: list[dict[str, str]] = []
    current: dict[str, str] | None = None
    field = "evidence"

    def flush():
        nonlocal current
        if current:
            for key in ("name", "status", "evidence", "development"):
                current[key] = _clean(current.get(key, ""))
            current["name"] = _display_name(current.get("name", ""))
            if any(current.get(key) for key in ("name", "status", "evidence", "development")):
                rows.append(current)
            current = None

    for raw in lines:
        line = _strip_markdown(raw, strip_bullet=True)
        if not line:
            continue

        match = STATUS_PATTERN.match(line)
        if match:
            flush()
            current = {
                "name": _display_name(match.group("name")),
                "status": match.group("status").upper(),
                "evidence": "",
                "development": "",
            }
            field = "evidence"
            continue

        if current is None:
            current = {
                "name": "Review narrative",
                "status": "",
                "evidence": "",
                "development": "",
            }

        lower = line.lower()
        if lower.startswith("evidence:"):
            field = "evidence"
            line = line.split(":", 1)[1].strip()
        elif lower.startswith("development note:"):
            field = "development"
            line = line.split(":", 1)[1].strip()
        elif lower.startswith("development opportunity:"):
            field = "development"
            line = line.split(":", 1)[1].strip()

        if line:
            current[field] = (current.get(field, "") + " " + line).strip()

    flush()
    return rows


def parse_competency_synthesis(lines: list[str]) -> list[dict[str, str]]:
    rows: list[dict[str, str]] = []
    current: dict[str, str] | None = None
    field = "evidence"

    def flush():
        nonlocal current
        if current:
            for key in ("name", "strength", "evidence", "development"):
                current[key] = _clean(current.get(key, ""))
            current["name"] = _display_name(current.get("name", ""))
            if any(current.get(key) for key in ("name", "strength", "evidence", "development")):
                rows.append(current)
            current = None

    for raw in lines:
        line = _strip_markdown(raw, strip_bullet=True)
        if not line:
            continue

        match = COMPETENCY_PATTERN.match(line)
        if match:
            flush()
            current = {
                "name": _display_name(match.group("name")),
                "strength": _clean(match.group("strength")),
                "evidence": "",
                "development": "",
            }
            field = "evidence"
            continue

        status_match = STATUS_PATTERN.match(line)
        if status_match and "competency" in status_match.group("name").lower():
            flush()
            current = {
                "name": _display_name(status_match.group("name")),
                "strength": status_match.group("status").upper(),
                "evidence": "",
                "development": "",
            }
            field = "evidence"
            continue

        if current is None:
            current = {
                "name": "Competency synthesis",
                "strength": "",
                "evidence": "",
                "development": "",
            }

        lower = line.lower()
        if lower.startswith("observed evidence:") or lower.startswith("evidence:"):
            field = "evidence"
            line = line.split(":", 1)[1].strip()
        elif lower.startswith("development opportunity:") or lower.startswith("development note:"):
            field = "development"
            line = line.split(":", 1)[1].strip()

        if line:
            current[field] = (current.get(field, "") + " " + line).strip()

    flush()
    return rows


def _set_document_defaults(doc: Document):
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = Inches(0.45)
    section.bottom_margin = Inches(0.45)
    section.left_margin = Inches(0.45)
    section.right_margin = Inches(0.45)

    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9)

    for name, size in (("Title", 18), ("Heading 1", 13), ("Heading 2", 11)):
        style = doc.styles[name]
        style.font.name = "Aptos Display" if name == "Title" else "Aptos"
        style.font.size = Pt(size)


def _add_heading(doc: Document, text: str, level: int = 1):
    paragraph = doc.add_heading(text, level=level)
    paragraph.paragraph_format.space_before = Pt(8)
    paragraph.paragraph_format.space_after = Pt(4)
    return paragraph


def _add_text_lines(doc: Document, lines: list[str]):
    for raw in lines:
        line = _strip_markdown(raw, strip_bullet=False)
        if not line:
            continue
        if line.startswith(("- ", "• ", "+ ")):
            paragraph = doc.add_paragraph(style="List Bullet")
            paragraph.add_run(line[2:].strip())
        else:
            paragraph = doc.add_paragraph()
            paragraph.add_run(line)
        paragraph.paragraph_format.space_after = Pt(3)


def _add_metrics_table(doc: Document, metrics):
    values = [
        ("Duration", _format_duration(metrics.duration_seconds)),
        ("Coach / Coachee word share", f"{metrics.coach_share_pct}% / {metrics.coachee_share_pct}%"),
        ("Coach turns", str(metrics.coach_turns)),
        ("Coachee turns", str(metrics.coachee_turns)),
        ("Coach questions", str(metrics.coach_questions)),
        ("Stacked-question turns", str(metrics.stacked_question_turns)),
        ("Average coach turn", f"{metrics.average_coach_words:.1f} words"),
        ("Longest coach turn", f"{metrics.longest_coach_turn_words} words"),
        ("Interruption notes", str(metrics.interruptions)),
    ]
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.rows[0].cells[0].text = "Session metric"
    table.rows[0].cells[1].text = "Value"
    for cell in table.rows[0].cells:
        for run in cell.paragraphs[0].runs:
            run.bold = True
    for label, value in values:
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = value
    return table


def _format_duration(seconds: int) -> str:
    total = max(0, int(seconds))
    hours, remainder = divmod(total, 3600)
    minutes, secs = divmod(remainder, 60)
    if hours:
        return f"{hours}:{minutes:02d}:{secs:02d}"
    return f"{minutes}:{secs:02d}"


def _timestamps(value: str) -> set[str]:
    return set(TIMESTAMP_PATTERN.findall(value or ""))


def _review_cell_text(row: dict[str, str]) -> str:
    parts = []
    name = row.get("name", "").strip()
    status = row.get("status", "").strip()
    evidence = row.get("evidence", "").strip()
    development = row.get("development", "").strip()

    if name:
        parts.append(name)
    if status:
        parts.append(status)
    if evidence:
        parts.append(evidence)
    if development:
        parts.append(f"Development: {development}")
    return "\n".join(parts)


def _add_annotated_transcript_table(doc: Document, transcript_rows, behavior_rows):
    """Create the review as an extension of the transcript itself."""
    visible_rows = [
        row for row in (transcript_rows or [])
        if len(row) >= 3 and row[1] in ("Coach", "Coachee")
    ]

    observations_by_stamp: dict[str, list[str]] = {}
    unmatched: list[dict[str, str]] = []
    for behavior in behavior_rows:
        stamps = _timestamps(behavior.get("evidence", ""))
        if not stamps:
            unmatched.append(behavior)
            continue
        review_text = _review_cell_text(behavior)
        for stamp in stamps:
            observations_by_stamp.setdefault(stamp, []).append(review_text)

    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    table.autofit = False
    headers = ("Speaker", "Timestamp", "Transcript", "Review / ICF Observation")
    widths = (Inches(0.9), Inches(0.9), Inches(4.0), Inches(4.15))

    for idx, (header, width) in enumerate(zip(headers, widths)):
        cell = table.rows[0].cells[idx]
        cell.width = width
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(header)
        run.bold = True
        run.font.size = Pt(8.5)

    for stamp, role, transcript in visible_rows:
        cells = table.add_row().cells
        values = (
            role,
            stamp,
            (transcript or "").strip(),
            "\n\n".join(observations_by_stamp.get(stamp, [])),
        )
        for idx, (value, width) in enumerate(zip(values, widths)):
            cells[idx].width = width
            paragraph = cells[idx].paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            run = paragraph.add_run(str(value))
            run.font.size = Pt(8.5)

    return table, unmatched


def _add_competency_table(doc: Document, rows: list[dict[str, str]]):
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    headers = ("Competency", "Evidence strength", "Observed evidence", "Development opportunity")
    widths = (Inches(1.75), Inches(1.4), Inches(3.35), Inches(3.4))
    for idx, (header, width) in enumerate(zip(headers, widths)):
        table.rows[0].cells[idx].width = width
        table.rows[0].cells[idx].text = header
        for run in table.rows[0].cells[idx].paragraphs[0].runs:
            run.bold = True
            run.font.size = Pt(8.5)

    for row in rows:
        cells = table.add_row().cells
        values = (
            row.get("name", ""),
            row.get("strength", ""),
            row.get("evidence", ""),
            row.get("development", ""),
        )
        for idx, (value, width) in enumerate(zip(values, widths)):
            cells[idx].width = width
            cells[idx].text = value
            for paragraph in cells[idx].paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.size = Pt(8.5)
    return table


def export_review_docx(
    review_text: str,
    filename: str | Path,
    level: str,
    metrics,
    scenario=None,
    transcript_rows=None,
):
    """Export a developmental review as an annotated transcript plus synthesis."""
    if not (review_text or "").strip():
        raise ValueError("There is no coaching review to export.")

    level = (level or "PCC").upper()
    sections = split_review_sections(review_text)
    behavior_lines = sections.get("MARKER / BEHAVIORAL EVIDENCE", [])
    competency_lines = sections.get("COMPETENCY SYNTHESIS", [])
    behavior_rows = parse_behavior_evidence(behavior_lines)
    competency_rows = parse_competency_synthesis(competency_lines)

    doc = Document()
    _set_document_defaults(doc)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Presence Coach - Practice Review")

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run(f"Developmental lens: {level}")
    run.bold = True
    run.font.size = Pt(11)

    basis = sections.get("ASSESSMENT BASIS", [])
    if basis:
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(f"Assessment basis: {_clean(' '.join(basis))}")
        run.font.size = Pt(8.5)

    if scenario:
        scenario_name = scenario.get("title", "Practice scenario")
        environment = scenario.get("environment", "")
        visible_problem = scenario.get("visible_problem", "")
        _add_heading(doc, "Practice scenario", 1)
        doc.add_paragraph(f"{scenario_name}" + (f" - {environment}" if environment else ""))
        if visible_problem:
            doc.add_paragraph(visible_problem)

    _add_heading(doc, "Session metrics", 1)
    _add_metrics_table(doc, metrics)
    note = doc.add_paragraph()
    note_run = note.add_run(
        "Speaking share is estimated from transcript word count, not measured audio time. "
        "Question counts depend on provider transcription punctuation."
    )
    note_run.italic = True
    note_run.font.size = Pt(8)

    _add_heading(doc, "Annotated Transcript", 1)
    intro = doc.add_paragraph()
    intro_run = intro.add_run(
        "The review is attached to the transcript moments that support each observation. "
        "Blank review cells simply mean no specific behavior-level observation was attached to that turn."
    )
    intro_run.font.size = Pt(8.5)
    intro_run.italic = True

    unmatched = behavior_rows
    if transcript_rows:
        _, unmatched = _add_annotated_transcript_table(doc, transcript_rows, behavior_rows)
    else:
        doc.add_paragraph(
            "Transcript rows were not available to this export. Behavioral evidence is listed below."
        )

    if unmatched:
        _add_heading(doc, "Behavioral observations without a matched transcript timestamp", 2)
        for row in unmatched:
            paragraph = doc.add_paragraph()
            paragraph.add_run(_review_cell_text(row))

    _add_heading(doc, "What the Coach Did Well", 1)
    _add_text_lines(doc, sections.get("WHAT THE COACH DID WELL", []))

    _add_heading(doc, "Competency Synthesis", 1)
    if competency_rows:
        _add_competency_table(doc, competency_rows)
    else:
        _add_text_lines(doc, competency_lines)

    for heading in (
        "PATTERNS TO WATCH",
        "THREE HIGH-LEVERAGE PRACTICE EDGES",
        "MOMENTS WORTH REVISITING",
        "BOTTOM LINE",
    ):
        _add_heading(doc, heading.title(), 1)
        _add_text_lines(doc, sections.get(heading, []))

    populated_sections = sum(
        1 for heading in SECTION_HEADINGS if any(_clean(x) for x in sections.get(heading, []))
    )
    if not behavior_rows or populated_sections < 4:
        _add_heading(doc, "Full Generated Review", 1)
        paragraph = doc.add_paragraph()
        paragraph.add_run(_strip_markdown(review_text, strip_bullet=False))

    disclaimer = doc.add_paragraph()
    disclaimer.paragraph_format.space_before = Pt(10)
    run = disclaimer.add_run(
        "Developmental AI review only - not an official ICF assessment, score, pass/fail result, "
        "or credential-readiness decision. Review important conclusions with a qualified human mentor coach."
    )
    run.bold = True
    run.font.size = Pt(8.5)

    doc.save(str(filename))
