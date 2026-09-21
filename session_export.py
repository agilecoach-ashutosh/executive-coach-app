"""Session export helpers for Presence Coach.

- Word transcript export with Speaker / Timestamp / Transcript columns.
- In-memory two-sided session audio capture and MP3 export.

Audio is not written to disk during the live session. It is retained in memory and
written only when the user explicitly chooses Export audio.
"""
from __future__ import annotations

import threading
from dataclasses import dataclass
from pathlib import Path

import lameenc
import numpy as np
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

TARGET_AUDIO_RATE = 24000
MP3_BITRATE_KBPS = 96
MAX_AUDIO_DURATION_SECONDS = 90 * 60
MAX_AUDIO_BYTES = 192 * 1024 * 1024


def format_elapsed(seconds: float) -> str:
    total = max(0, int(seconds))
    hours, remainder = divmod(total, 3600)
    minutes, secs = divmod(remainder, 60)
    return f"{hours:02d}:{minutes:02d}:{secs:02d}"


def export_transcript_docx(rows, filename: str | Path, title: str = "Presence Coach — Session Transcript"):
    """Write the visible coaching transcript as a three-column Word table."""
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.6)
    section.right_margin = Inches(0.6)

    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10)

    heading = doc.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = heading.add_run(title)
    run.bold = True
    run.font.size = Pt(15)

    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note_run = note.add_run(
        "Timestamps show elapsed time from the start of the coaching session. "
        "Speech transcription may contain recognition errors."
    )
    note_run.font.size = Pt(8)

    visible_rows = [row for row in rows if len(row) >= 3 and row[1] in ("Coach", "Coachee")]
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    table.autofit = False

    widths = (Inches(1.35), Inches(1.25), Inches(4.7))
    headers = ("Speaker (Coach / Coachee)", "Timestamp", "Transcript")
    header_cells = table.rows[0].cells
    for idx, (text, width) in enumerate(zip(headers, widths)):
        header_cells[idx].width = width
        paragraph = header_cells[idx].paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        header_run = paragraph.add_run(text)
        header_run.bold = True
        header_run.font.size = Pt(9)

    for stamp, role, text in visible_rows:
        cells = table.add_row().cells
        values = (role, stamp, (text or "").strip())
        for idx, (value, width) in enumerate(zip(values, widths)):
            cells[idx].width = width
            paragraph = cells[idx].paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.add_run(str(value))

    doc.save(str(filename))


@dataclass
class _AudioSegment:
    start_offset: float
    sample_rate: int
    channels: int
    data: bytearray

    @property
    def duration(self) -> float:
        frames = len(self.data) / max(1, self.channels * 2)
        return frames / max(1, self.sample_rate)


class SessionAudioRecorder:
    """Thread-safe in-memory recorder for the human and AI audio streams.

    Presence is turn-based, so the two streams are normally non-overlapping. Segments
    are time-aligned to the same monotonic session start and encoded to one mono MP3.
    """

    def __init__(self):
        self.started_at: float | None = None
        self._segments: list[_AudioSegment] = []
        self._lock = threading.Lock()
        self._captured_bytes = 0
        self._truncated = False

    def start(self, started_at: float):
        with self._lock:
            self.started_at = float(started_at)
            self._segments.clear()
            self._captured_bytes = 0
            self._truncated = False

    def clear(self):
        with self._lock:
            self.started_at = None
            self._segments.clear()
            self._captured_bytes = 0
            self._truncated = False

    def append(self, pcm: bytes, sample_rate: int, channels: int = 1, now: float | None = None):
        if not pcm or self.started_at is None:
            return
        import time

        timestamp = time.monotonic() if now is None else float(now)
        offset = max(0.0, timestamp - self.started_at)
        incoming = bytes(pcm)
        frame_size = max(1, int(channels) * 2)

        with self._lock:
            remaining_duration_bytes = max(
                0,
                int((MAX_AUDIO_DURATION_SECONDS - offset) * sample_rate) * frame_size,
            )
            remaining_memory_bytes = max(0, MAX_AUDIO_BYTES - self._captured_bytes)
            allowed = min(len(incoming), remaining_duration_bytes, remaining_memory_bytes)
            allowed -= allowed % frame_size
            if allowed < len(incoming):
                self._truncated = True
            incoming = incoming[:allowed]
            if not incoming:
                return

            # Merge short contiguous callback chunks so long sessions do not create
            # tens of thousands of tiny Python objects. Cap merged blocks at 5 sec.
            if self._segments:
                last = self._segments[-1]
                expected = last.start_offset + last.duration
                if (
                    last.sample_rate == sample_rate
                    and last.channels == channels
                    and abs(offset - expected) <= 0.12
                    and last.duration < 5.0
                ):
                    last.data.extend(incoming)
                    self._captured_bytes += len(incoming)
                    return
            self._segments.append(
                _AudioSegment(offset, int(sample_rate), int(channels), bytearray(incoming))
            )
            self._captured_bytes += len(incoming)

    def has_audio(self) -> bool:
        with self._lock:
            return any(segment.data for segment in self._segments)

    def is_truncated(self) -> bool:
        with self._lock:
            return self._truncated

    def duration_seconds(self) -> float:
        with self._lock:
            if not self._segments:
                return 0.0
            return max(segment.start_offset + segment.duration for segment in self._segments)

    @staticmethod
    def _mono_int16(segment: _AudioSegment) -> np.ndarray:
        values = np.frombuffer(bytes(segment.data), dtype="<i2")
        if segment.channels > 1:
            usable = len(values) - (len(values) % segment.channels)
            values = values[:usable].reshape(-1, segment.channels).astype(np.int32).mean(axis=1)
            values = np.clip(values, -32768, 32767).astype(np.int16)
        return values.astype(np.int16, copy=False)

    @staticmethod
    def _resample(values: np.ndarray, source_rate: int, target_rate: int) -> np.ndarray:
        if not len(values) or source_rate == target_rate:
            return values.astype(np.int16, copy=False)
        target_len = max(1, round(len(values) * target_rate / source_rate))
        source_positions = np.arange(len(values), dtype=np.float64)
        target_positions = np.linspace(0, max(0, len(values) - 1), target_len)
        result = np.interp(target_positions, source_positions, values.astype(np.float64))
        return np.clip(np.rint(result), -32768, 32767).astype("<i2")

    @staticmethod
    def _encode_silence(encoder, file_handle, samples: int):
        block = np.zeros(TARGET_AUDIO_RATE, dtype="<i2").tobytes()
        remaining = max(0, int(samples))
        while remaining:
            take = min(remaining, TARGET_AUDIO_RATE)
            encoded = encoder.encode(block[: take * 2])
            if encoded:
                file_handle.write(encoded)
            remaining -= take

    def export_mp3(self, filename: str | Path):
        with self._lock:
            # Export is offered after the engine stops. A shallow snapshot avoids
            # duplicating the complete recording in memory before encoding begins.
            segments = [s for s in self._segments if s.data]

        if not segments:
            raise RuntimeError("No session audio is available to export.")

        segments.sort(key=lambda item: item.start_offset)
        encoder = lameenc.Encoder()
        encoder.set_bit_rate(MP3_BITRATE_KBPS)
        encoder.set_in_sample_rate(TARGET_AUDIO_RATE)
        encoder.set_channels(1)
        encoder.set_quality(2)

        cursor = 0
        output_path = Path(filename)
        with output_path.open("wb") as out:
            for segment in segments:
                values = self._mono_int16(segment)
                values = self._resample(values, segment.sample_rate, TARGET_AUDIO_RATE)
                if not len(values):
                    continue

                start_sample = max(0, round(segment.start_offset * TARGET_AUDIO_RATE))
                if start_sample > cursor:
                    self._encode_silence(encoder, out, start_sample - cursor)
                    cursor = start_sample
                elif start_sample < cursor:
                    overlap = cursor - start_sample
                    if overlap >= len(values):
                        continue
                    values = values[overlap:]

                encoded = encoder.encode(values.astype("<i2", copy=False).tobytes())
                if encoded:
                    out.write(encoded)
                cursor += len(values)

            tail = encoder.flush()
            if tail:
                out.write(tail)

        if not output_path.exists() or output_path.stat().st_size == 0:
            raise RuntimeError("MP3 export did not produce an audio file.")
