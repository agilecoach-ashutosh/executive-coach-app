import tempfile
import unittest
from pathlib import Path

from docx import Document

from review_export import (
    export_review_docx,
    parse_behavior_evidence,
    parse_competency_synthesis,
    split_review_sections,
)
from reviewer import SessionMetrics


SAMPLE_REVIEW = """DEVELOPMENTAL REVIEW — PCC
ASSESSMENT BASIS: ICF PCC Minimum Skills Requirements, effective January 1, 2026

WHAT THE COACH DID WELL
- The coach stayed with the client's agenda at [00:00:18].

MARKER / BEHAVIORAL EVIDENCE
Competency 3 — Clear client-owned session outcome — OBSERVED
Evidence: [00:00:18] The coach asked what would make the conversation useful.
Development note: Reconfirm the outcome when the topic shifts.

Competency 5 — Creates reflective space — PARTIAL EVIDENCE
Evidence: [00:03:10] A pause was visible in the transcript flow, but audio nuance is unavailable.
Development note: Allow more space after emotionally significant statements.

COMPETENCY SYNTHESIS
Establishes and Maintains Agreements — Evidence strength: Strong
Observed evidence: The outcome was client-owned and revisited.
Development opportunity: Re-contract explicitly after a major shift.

PATTERNS TO WATCH
- Avoid stacked questions.

THREE HIGH-LEVERAGE PRACTICE EDGES
- Ask one question at a time.

MOMENTS WORTH REVISITING
- [00:02:20] One coach turn contained two questions.

BOTTOM LINE
The session shows a client-centered foundation. Developmental AI review only — not an official ICF assessment.
"""


MARKDOWN_REVIEW = """## DEVELOPMENTAL REVIEW — ACC
**ASSESSMENT BASIS:** ICF ACC Minimum Skills Requirements, effective January 1, 2026

### WHAT THE COACH DID WELL
1. **Asked Open Inquiry Early (Competency 7):** At [00:00:34], the coach asked a clean open question.

### MARKER / BEHAVIORAL EVIDENCE
***[Competency 1 — Stays in Coaching Role / Avoids Directing] — INCONSISTENT***
* **Evidence:** At [00:01:56], the coach gave explicit advice.
* **Development note:** Return ownership to the client and avoid prescribing tactics.

***[Competency 3 — Explores Client's Chosen Topic] — OBSERVED***
* **Evidence:** At [00:00:34] and [00:00:59], the coach explored the presenting topic.

### COMPETENCY SYNTHESIS
**Competency 3 — Establishes and Maintains Agreements — Evidence strength: Developing**
**Observed evidence:** The coach explored the topic but did not establish a clear session outcome.
**Development opportunity:** Partner explicitly on what the client wants from the session.

### PATTERNS TO WATCH
- Advice-giving replaced coaching at [00:01:56].

### THREE HIGH-LEVERAGE PRACTICE EDGES
- Establish a client-owned session outcome before deep exploration.

### MOMENTS WORTH REVISITING
- [00:01:56] The coach prescribed an escalation strategy.

### BOTTOM LINE
The session contains some useful inquiry but also a clear move out of the coaching role.
"""


ACC_RATING_REVIEW = """MARKER / BEHAVIORAL EVIDENCE
A7.3 — Asks clear, open-ended questions, one at a time — MEETS THE STANDARD
Evidence: [00:00:34] The coach asks one clear open question.
Development note: Continue keeping questions concise.
"""


def _metrics():
    return SessionMetrics(
        duration_seconds=240,
        coach_turns=8,
        coachee_turns=9,
        coach_words=120,
        coachee_words=310,
        coach_share_pct=28,
        coachee_share_pct=72,
        coach_questions=7,
        stacked_question_turns=1,
        average_coach_words=15.0,
        longest_coach_turn_words=28,
        interruptions=0,
    )


SAMPLE_TRANSCRIPT = [
    ("00:00:18", "Coach", "What would make this conversation useful for you today?"),
    ("00:00:29", "Coachee", "I want to work out how to handle a difficult conversation."),
    ("00:02:20", "Coach", "What have you tried, and what else could you do?"),
    ("00:03:10", "Coach", "What are you noticing now?"),
]


MARKDOWN_TRANSCRIPT = [
    ("00:00:34", "Coach", "What would you like to explore about this?"),
    ("00:00:59", "Coach", "What feels most important in that?"),
    ("00:01:20", "Coachee", "I am worried about how my manager will react."),
    ("00:01:56", "Coach", "You should escalate this to your director."),
]


class ReviewExportTests(unittest.TestCase):
    def test_sections_and_evidence_parse(self):
        sections = split_review_sections(SAMPLE_REVIEW)
        behavior = parse_behavior_evidence(sections["MARKER / BEHAVIORAL EVIDENCE"])
        synthesis = parse_competency_synthesis(sections["COMPETENCY SYNTHESIS"])

        self.assertEqual(len(behavior), 2)
        self.assertEqual(behavior[0]["status"], "OBSERVED")
        self.assertIn("00:00:18", behavior[0]["evidence"])
        self.assertEqual(len(synthesis), 1)
        self.assertEqual(synthesis[0]["strength"], "Strong")

    def test_acc_observation_rating_parses(self):
        sections = split_review_sections(ACC_RATING_REVIEW)
        behavior = parse_behavior_evidence(sections["MARKER / BEHAVIORAL EVIDENCE"])
        self.assertEqual(len(behavior), 1)
        self.assertEqual(behavior[0]["status"], "MEETS THE STANDARD")
        self.assertIn("00:00:34", behavior[0]["evidence"])

    def test_markdown_review_from_provider_parses_into_sections_and_rows(self):
        sections = split_review_sections(MARKDOWN_REVIEW)
        self.assertIn("Asked Open Inquiry Early", " ".join(sections["WHAT THE COACH DID WELL"]))
        self.assertIn("Advice-giving", " ".join(sections["PATTERNS TO WATCH"]))
        self.assertIn("The session contains", " ".join(sections["BOTTOM LINE"]))

        behavior = parse_behavior_evidence(sections["MARKER / BEHAVIORAL EVIDENCE"])
        self.assertEqual(len(behavior), 2)
        self.assertEqual(behavior[0]["status"], "INCONSISTENT")
        self.assertIn("Stays in Coaching Role", behavior[0]["name"])
        self.assertIn("00:01:56", behavior[0]["evidence"])
        self.assertIn("avoid prescribing", behavior[0]["development"])

        synthesis = parse_competency_synthesis(sections["COMPETENCY SYNTHESIS"])
        self.assertEqual(len(synthesis), 1)
        self.assertEqual(synthesis[0]["strength"], "Developing")
        self.assertIn("clear session outcome", synthesis[0]["evidence"])

    def test_word_review_contains_annotated_transcript(self):
        scenario = {
            "title": "Leadership transition",
            "environment": "Technology",
            "visible_problem": "A new manager is struggling to delegate.",
        }

        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "review.docx"
            export_review_docx(
                SAMPLE_REVIEW,
                path,
                "PCC",
                _metrics(),
                scenario,
                SAMPLE_TRANSCRIPT,
            )
            doc = Document(path)

            self.assertGreaterEqual(len(doc.tables), 3)
            metric_headers = [cell.text for cell in doc.tables[0].rows[0].cells]
            transcript_headers = [cell.text for cell in doc.tables[1].rows[0].cells]
            competency_headers = [cell.text for cell in doc.tables[2].rows[0].cells]

            self.assertEqual(metric_headers, ["Session metric", "Value"])
            self.assertEqual(
                transcript_headers,
                ["Speaker", "Timestamp", "Transcript", "Review / ICF Observation"],
            )
            self.assertEqual(
                competency_headers,
                ["Competency", "Evidence strength", "Observed evidence", "Development opportunity"],
            )

            first_review_row = [cell.text for cell in doc.tables[1].rows[1].cells]
            self.assertEqual(first_review_row[0], "Coach")
            self.assertEqual(first_review_row[1], "00:00:18")
            self.assertIn("What would make this conversation useful", first_review_row[2])
            self.assertIn("Clear client-owned session outcome", first_review_row[3])
            self.assertIn("OBSERVED", first_review_row[3])

            coachee_row = [cell.text for cell in doc.tables[1].rows[2].cells]
            self.assertEqual(coachee_row[0], "Coachee")
            self.assertEqual(coachee_row[3], "")

    def test_markdown_word_review_attaches_observation_to_matching_timestamp(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "markdown-review.docx"
            export_review_docx(
                MARKDOWN_REVIEW,
                path,
                "ACC",
                _metrics(),
                transcript_rows=MARKDOWN_TRANSCRIPT,
            )
            doc = Document(path)

            self.assertGreaterEqual(len(doc.tables), 3)
            transcript_table = doc.tables[1]
            rows_by_stamp = {
                row.cells[1].text: [cell.text for cell in row.cells]
                for row in transcript_table.rows[1:]
            }

            self.assertIn("INCONSISTENT", rows_by_stamp["00:01:56"][3])
            self.assertIn("Stays in Coaching Role", rows_by_stamp["00:01:56"][3])
            self.assertIn("OBSERVED", rows_by_stamp["00:00:34"][3])
            self.assertIn("Explores Client's Chosen Topic", rows_by_stamp["00:00:59"][3])

            competency_text = " ".join(
                cell.text for row in doc.tables[2].rows for cell in row.cells
            )
            paragraph_text = " ".join(paragraph.text for paragraph in doc.paragraphs)

            self.assertIn("Developing", competency_text)
            self.assertIn("Advice-giving", paragraph_text)
            self.assertIn("The session contains", paragraph_text)


if __name__ == "__main__":
    unittest.main()
