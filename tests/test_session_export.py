import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import numpy as np
from docx import Document

from session_export import SessionAudioRecorder, export_transcript_docx, format_elapsed


class SessionExportTests(unittest.TestCase):
    def test_audio_memory_cap_sets_truncation_and_clear_resets_it(self):
        recorder = SessionAudioRecorder()
        recorder.start(10.0)
        with patch("session_export.MAX_AUDIO_BYTES", 8):
            recorder.append(b"\x01\x00" * 6, 16_000, now=10.1)

        self.assertTrue(recorder.has_audio())
        self.assertTrue(recorder.is_truncated())
        self.assertEqual(sum(len(segment.data) for segment in recorder._segments), 8)

        recorder.clear()
        self.assertFalse(recorder.has_audio())
        self.assertFalse(recorder.is_truncated())

    def test_audio_duration_cap_rejects_late_frames(self):
        recorder = SessionAudioRecorder()
        recorder.start(10.0)
        with patch("session_export.MAX_AUDIO_DURATION_SECONDS", 1):
            recorder.append(b"\x01\x00" * 10, 16_000, now=11.1)

        self.assertFalse(recorder.has_audio())
        self.assertTrue(recorder.is_truncated())

    def test_elapsed_timestamp_format(self):
        self.assertEqual(format_elapsed(0), "00:00:00")
        self.assertEqual(format_elapsed(137), "00:02:17")
        self.assertEqual(format_elapsed(3661), "01:01:01")

    def test_word_transcript_has_requested_three_columns(self):
        rows = [
            ("00:00:04", "Coach", "What would make this useful today?"),
            ("00:00:12", "Coachee", "I want to think about delegation."),
            ("00:00:15", "Session note", "Internal note should not be a speaker row."),
        ]
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "transcript.docx"
            export_transcript_docx(rows, path)
            doc = Document(path)
            self.assertEqual(len(doc.tables), 1)
            table = doc.tables[0]
            self.assertEqual(
                [cell.text for cell in table.rows[0].cells],
                ["Speaker (Coach / Coachee)", "Timestamp", "Transcript"],
            )
            # Transcript rows are stored internally as (timestamp, role, text),
            # but the exported Word document is intentionally Speaker, Timestamp, Transcript.
            self.assertEqual(
                [cell.text for cell in table.rows[1].cells],
                [rows[0][1], rows[0][0], rows[0][2]],
            )
            self.assertEqual(
                [cell.text for cell in table.rows[2].cells],
                [rows[1][1], rows[1][0], rows[1][2]],
            )
            self.assertEqual(len(table.rows), 3)

    def test_mp3_export_writes_audio(self):
        recorder = SessionAudioRecorder()
        recorder.start(100.0)
        half_second = np.zeros(8000, dtype="<i2").tobytes()
        recorder.append(half_second, 16000, 1, now=100.0)
        recorder.append(half_second, 16000, 1, now=100.5)
        self.assertTrue(recorder.has_audio())

        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "session.mp3"
            recorder.export_mp3(path)
            self.assertTrue(path.exists())
            self.assertGreater(path.stat().st_size, 0)


if __name__ == "__main__":
    unittest.main()
